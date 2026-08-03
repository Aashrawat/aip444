"""Generate Lab 09 submission Word document."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(r"C:\Users\aashr\OneDrive\Desktop\AIP444")
LAB = ROOT / "labs" / "lab-09"
OUT = ROOT / "lab-09.docx"


def set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def add_code(doc: Document, text: str) -> None:
    for line in text.splitlines() or [""]:
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def main() -> None:
    doc = Document()
    set_normal_style(doc)

    title = doc.add_heading("Lab 09: Source Credibility Analyzer Agent", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "Aashrawat Shrestha — 179413232")
    add_para(doc, "AIP444 — OpenAI Agents SDK (TypeScript) + OpenRouter + Tavily")
    add_para(
        doc,
        "Repository: https://github.com/Aashrawat/aip444/tree/main/labs/lab-09",
    )

    # ----- 1. GitHub -----
    add_heading(doc, "1. GitHub Repository", 1)
    add_para(
        doc,
        "Code link: https://github.com/Aashrawat/aip444/tree/main/labs/lab-09",
    )
    add_para(
        doc,
        "The README.md in that folder explains installation and API key configuration. Summary:",
    )
    add_heading(doc, "Install dependencies", 2)
    add_code(
        doc,
        "cd labs/lab-09\nnpm install",
    )
    add_heading(doc, "Configure API keys", 2)
    add_para(
        doc,
        "Add keys to the repository root .env (one level above labs/):",
    )
    add_code(
        doc,
        "OPENROUTER_API_KEY=sk-or-...\n"
        "TAVILY_API_KEY=tvly-...\n"
        "# Optional — defaults to openai/gpt-5.4-mini (skips flaky openrouter/free)\n"
        "OPENROUTER_MODEL=openai/gpt-5.4-mini\n"
        "MAX_TURNS=20",
    )
    add_heading(doc, "Run", 2)
    add_code(
        doc,
        'npm run analyze -- "https://example.com/article"',
    )
    add_para(
        doc,
        "Tool logs and the agent trace print to stderr; the final message prints to stdout. "
        "Reports are written under labs/lab-09/reports/.",
    )
    add_para(
        doc,
        "Packages: @openai/agents, openai (OpenRouter), @tavily/core, zod, dotenv.",
    )

    # ----- 2. System prompt -----
    add_heading(doc, "2. System Prompt", 1)
    add_para(
        doc,
        "Full system prompt from labs/lab-09/SYSTEM_PROMPT.md:",
    )
    add_code(doc, read(LAB / "SYSTEM_PROMPT.md"))

    # ----- 3. Schema -----
    add_heading(doc, "3. Schema for assess_credibility (Zod)", 1)
    add_para(
        doc,
        "The assess_credibility tool uses a Zod rubric (schemas.ts). "
        "The implementation returns the evaluation unaltered (think-tool pattern).",
    )
    add_code(doc, read(LAB / "schemas.ts"))
    add_para(doc, "Tool wiring excerpt from tools.ts:")
    add_code(
        doc,
        """export const assessCredibilityTool = tool({
  name: 'assess_credibility',
  description:
    'Record a structured credibility evaluation using the full rubric. Call this AFTER investigation...',
  parameters: CredibilityEvaluationSchema,
  async execute(evaluation) {
    return { status: 'evaluation_recorded', evaluation };
  },
});""",
    )

    # ----- 4. Trace -----
    add_heading(doc, "4. Agent Trace (BBC test run)", 1)
    add_para(
        doc,
        "Source evaluated: https://www.bbc.com/news/business-65086798 "
        "(“Friend or foe: Can computer coders trust ChatGPT?”).",
    )
    add_para(
        doc,
        "Model: openai/gpt-5.4-mini via OpenRouter CustomModelProvider. "
        "maxTurns: 20. The trace below shows every tool call/result from result.newItems.",
    )
    trace_path = LAB / "traces" / "bbc-agent-trace-clean.txt"
    if not trace_path.exists():
        trace_path = LAB / "traces" / "bbc-full-trace.txt"
    trace_text = read(trace_path)
    # Keep document readable but still "full" — include entire cleaned trace
    add_code(doc, trace_text)

    add_heading(doc, "Trace summary", 2)
    add_para(
        doc,
        "Typical sequence on this run: read_url (article) → read_url (BBC About) → "
        "web_search (author credentials) → web_search (BBC editorial standards) → "
        "web_search (claim corroboration / ChatGPT coding security) → "
        "additional searches/reads → assess_credibility (overall: high) → save_report → final message.",
    )

    # ----- 5. Reports -----
    add_heading(doc, "5. Credibility Reports", 1)
    add_para(
        doc,
        "Two representative final Markdown reports produced by the agent:",
    )

    add_heading(doc, "5.1 High credibility — BBC news", 2)
    add_code(doc, read(LAB / "reports" / "bbc-business-65086798-credibility.md"))

    add_heading(doc, "5.2 Tricky / very low — Natural News", 2)
    add_code(
        doc,
        read(LAB / "reports" / "naturalnews-chatgpt-threat-to-humanity-credibility.md"),
    )

    add_heading(doc, "Additional reports generated during testing", 2)
    add_para(
        doc,
        "Also available in labs/lab-09/reports/: NIST (high), GitHub Blog Copilot research (medium), "
        "Medium personal opinion (low), plus research-project sources Peng et al. arXiv (high with caveats) "
        "and Anthropic Claude Code docs (high as official product docs).",
    )

    # ----- 6. Reflection -----
    add_heading(doc, "6. Reflection", 1)

    add_heading(doc, "How many steps? Logical process or detours?", 2)
    add_para(
        doc,
        "Across successful runs the agent typically used about 8–13 tool calls "
        "(often ~10–12 “steps” in the printed trace, counting calls and the final message). "
        "A common pattern was: read the source → search author → search publication → "
        "search claims / corroboration → optionally read About or linked pages → "
        "assess_credibility → save_report.",
    )
    add_para(
        doc,
        "Overall the process was logical and matched the system prompt. "
        "Some detours appeared: overlapping web_search queries on the same topic, "
        "or re-reading list/index pages that added little. Those rarely changed the verdict, "
        "but they burned turns. The strongest runs stayed close to the intended investigative order "
        "and finished well under the maxTurns=20 guardrail.",
    )

    add_heading(doc, "Effect of assess_credibility (think-tool ablation)", 2)
    add_para(
        doc,
        "I re-ran the same BBC URL with SKIP_ASSESS=1 (tools = read_url, web_search, save_report only). "
        "Both runs still rated the source high, but quality differed:",
    )
    add_para(
        doc,
        "With assess_credibility: the report includes a complete structured rubric "
        "(source_type enums, boolean claim flags, transparency_score, editorial_process, etc.) "
        "that mirrors the investigation. Fields like sources_cited=false vs corroborated=true "
        "force useful distinctions.",
    )
    add_para(
        doc,
        "Without assess_credibility: the report is still readable and reaches a sensible verdict, "
        "but the “structured evaluation” section becomes looser prose (“Moderate to strong,” "
        "“Low to moderate”) instead of schema-constrained fields. It is easier for the model to "
        "skip a dimension or blur “unverified anecdote” vs “corroborated claim.”",
    )
    add_para(
        doc,
        "This matches the Anthropic “think” tool idea: the power is in the schema forcing a "
        "complete intermediate judgment into context before the final write. The tool does no I/O; "
        "it improves discipline.",
    )

    add_heading(doc, "Unknown author / publication — hallucination vs routing", 2)
    add_para(
        doc,
        "When authorship was missing (NIST news release; Anthropic docs; Natural News 404 page), "
        "the agent generally followed the prompt: it checked About pages, searched, and recorded "
        "Unknown or institutional authorship rather than inventing a byline. "
        "On the broken NIST July URL, it discovered the live January page and documented the 404 "
        "instead of fabricating article text.",
    )
    add_para(
        doc,
        "I did revise the system prompt / schema after early tests. Polished Medium essays hosted "
        "under Medium “publications” (e.g., CodeX) were sometimes rated medium until I added "
        "explicit guidance that Medium publications are usually still self-published personal blogs "
        "unless a real editorial process is found. I also removed Zod .url() formats after the "
        "provider rejected format:uri in tool schemas.",
    )
    add_para(
        doc,
        "Residual risk: the agent can still mis-classify source_type (e.g., labeling an arXiv "
        "preprint as peer_reviewed_journal while correctly saying editorial_process=self_published). "
        "That is inconsistency inside one evaluation, not fabricating an author name, but it shows "
        "schema discipline is incomplete without stronger cross-field constraints.",
    )

    add_heading(doc, "Hardest source to evaluate correctly", 2)
    add_para(
        doc,
        "The hardest cases were (1) Natural News when the article URL returned a not-found page — "
        "the agent correctly rated very_low based on outlet reputation, but could not verify the "
        "specific article claims; and (2) vendor/research-borderline sources such as GitHub’s own "
        "Copilot productivity write-ups and Anthropic product docs, which look authoritative and "
        "often deserve high trust for “what the product claims,” but medium trust for comparative "
        "or evaluative claims because of incentive bias.",
    )
    add_para(
        doc,
        "Improvements I would make: add a dedicated reputation lookup tool (known-outlet database), "
        "stricter schema refinements (forbid peer_reviewed_journal unless editorial_process is "
        "peer_reviewed), allow longer read_url truncation for references, and prompt the agent to "
        "separate “trust for product description” vs “trust for independent evaluation.”",
    )

    add_heading(doc, "Agent vs evaluating sources myself", 2)
    add_para(
        doc,
        "Strengths: speed, consistent checklist coverage, parallel searching for author/outlet/"
        "corroboration, and producing a reusable Markdown artifact. It is good at catching "
        "obvious red flags (unknown authorship, misinformation outlets, broken URLs) and at "
        "pulling About/editorial pages I might skip when rushing.",
    )
    add_para(
        doc,
        "Weaknesses: shallow reading of long PDFs, occasional overconfidence on vendor docs, "
        "duplicate searches, blocked major-news fetches (e.g., Reuters), and imperfect "
        "classification. It cannot replace human judgment about whether a study’s methodology "
        "supports the claim I want to make in a paper.",
    )
    add_para(
        doc,
        "For my research project I would trust it as a first-pass triage and documentation aid, "
        "not as the final authority. I would still open primary papers, verify DOIs, and decide "
        "myself whether a source is citable for a specific claim.",
    )

    add_heading(doc, "Research-project sources (CLI vs IDE AI coding tools)", 2)
    add_para(
        doc,
        "I ran the agent on sources used in my CLI vs IDE research report:",
    )
    add_para(
        doc,
        "1) Peng et al. (2023), The Impact of AI on Developer Productivity: Evidence from GitHub Copilot — "
        "https://arxiv.org/abs/2302.06590 → agent verdict: high. "
        "This is appropriate as a primary experimental preprint for the 55.8% speed claim, with important "
        "caveats the agent noted: arXiv is not peer review, and some authors have GitHub/Microsoft ties. "
        "I still treat it as a credible core source for productivity evidence, but I would not overgeneralize "
        "beyond the study’s task setup.",
        bold=False,
    )
    add_para(
        doc,
        "2) Anthropic Claude Code overview docs — https://code.claude.com/docs/en/overview → agent verdict: high. "
        "Credible as official product documentation for describing what Claude Code is and how it works. "
        "For comparative claims (CLI better than IDE, safety superiority, etc.) I would rate it medium and "
        "pair it with independent research (e.g., Perry et al. 2023 on insecure code with AI assistants; "
        "Klemmer et al. 2024 on security practices).",
    )
    add_para(
        doc,
        "Overall credibility of my current research source mix: strong. The literature review leans on "
        "peer-reviewed / ACM CCS papers (Perry, Klemmer), established preprints with transparent methods "
        "(Peng; Ziegler et al. in CACM), and official docs (Anthropic, Cursor, Aider) clearly labeled as "
        "vendor documentation. I am not relying on low-credibility blogs for core claims. The agent’s "
        "BBC (high) and Natural News (very_low) tests align with how I already filter sources for the project.",
    )

    add_heading(doc, "Conclusion", 2)
    add_para(
        doc,
        "Lab 09 demonstrates a working multi-tool agent loop: read → search → structured think → report. "
        "The assess_credibility think tool measurably improves completeness of the final rubric. "
        "With prompt revisions for missing-info handling, the agent mostly avoids hallucinating authors "
        "and is useful as a research assistant for triage—while human verification remains essential "
        "for citation decisions.",
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
